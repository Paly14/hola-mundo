#!/usr/bin/env python3
"""
Convierte los trackers de Excel de Alpha Ecommerce en el archivo de datos
inicial de Alpha CRM (app/assets/js/datos-alpha.js).

Uso:
    python3 tools/importar-trackers.py CRM_Gabo.xlsx CRM_Setter.xlsx GESTION_ALUMNOS.xlsx

Los archivos se reconocen por su contenido, no por el nombre.
Requiere: pip install openpyxl
"""
import sys, json, hashlib, datetime, unicodedata, re

try:
    import openpyxl
except ImportError:
    sys.exit('Falta openpyxl:  pip install openpyxl')

MESES = {'ene': 1, 'feb': 2, 'mar': 3, 'abr': 4, 'may': 5, 'jun': 6, 'jul': 7,
         'ago': 8, 'sep': 9, 'oct': 10, 'nov': 11, 'dic': 12,
         'enero': 1, 'febrero': 2, 'marzo': 3, 'abril': 4, 'mayo': 5, 'junio': 6,
         'julio': 7, 'agosto': 8, 'septiembre': 9, 'octubre': 10,
         'noviembre': 11, 'diciembre': 12}


def sha256(text):
    return hashlib.sha256(text.encode('utf-8')).hexdigest()


def sinacentos(text):
    return ''.join(c for c in unicodedata.normalize('NFD', str(text))
                   if unicodedata.category(c) != 'Mn').lower().strip()


def iso(value):
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    if isinstance(value, str):
        m = re.match(r'^(\d{1,2})/(\d{1,2})/(\d{4})', value.strip())
        if m:  # dd/mm/aaaa, como se escribe en Argentina
            return '%s-%02d-%02d' % (m.group(3), int(m.group(2)), int(m.group(1)))
    return ''


def fecha_de_hoja(value, mes_hoja):
    """
    Excel interpreta "1/9/2026" como 9 de enero (formato de EE.UU.), pero en la
    hoja de septiembre significa 1 de septiembre. Cuando la fecha cae fuera del
    mes de la hoja, se dan vuelta día y mes.
    """
    if isinstance(value, datetime.datetime) and mes_hoja and value.month != mes_hoja:
        if value.month <= 31:
            return '%d-%02d-%02d' % (value.year, mes_hoja, value.month)
    return iso(value)


def telefono(value):
    digitos = re.sub(r'\D', '', str(value or '').split('.')[0])
    if not digitos:
        return ''
    if digitos.startswith('54'):
        return '+' + digitos
    return '+54 9 ' + digitos


def num(value):
    try:
        n = float(value)
        return int(n) if n == int(n) else round(n, 2)
    except (TypeError, ValueError):
        return None


def texto(value):
    if value is None:
        return ''
    return re.sub(r'\s+', ' ', str(value)).strip()


# ------------------------------------------------------------------ leads

SITUACIONES = {
    'seguimiento': 'Seguimiento',
    'adentro en seguimiento': 'Ganado',
    'esperando pago': 'Esperando pago',
    'cerrado': 'Ganado',
    'adentro': 'Ganado',
    'no califica': 'Perdido',
    'no calificaba': 'Perdido',
    'no se presento': 'No Show',
    'reagenda': 'Agendado',
    'perdido': 'Perdido',
}


def estado_de(situacion, presento):
    clave = sinacentos(situacion)
    if clave in SITUACIONES:
        return SITUACIONES[clave]
    if sinacentos(presento) == 'no':
        return 'No Show'
    return 'Presentado' if clave else 'Agendado'


def leer_closer(path, salida, config):
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        for fila in ws.iter_rows(min_row=1, max_row=3, values_only=True):
            for celda in fila:
                m = re.search(r'KPI MENSUAL:\s*\$?([\d.,]+)', str(celda or ''))
                if m:
                    monto = num(m.group(1).replace('.', '').replace(',', '.'))
                    if monto:
                        config['meta_cash'] = max(config.get('meta_cash', 0), monto)
    for ws in wb.worksheets:
        mes = MESES.get(sinacentos(ws.title).strip())
        if not mes:
            continue
        cabeceras = [texto(c.value) for c in ws[1]]
        indice = {sinacentos(h): i for i, h in enumerate(cabeceras) if h}

        def col(*claves):
            for clave in claves:
                for h, i in indice.items():
                    if h.startswith(sinacentos(clave)):
                        return i
            return None

        c_nombre = col('nombre del lead')
        if c_nombre is None:
            continue
        cols = {
            'tel': col('numero de contacto'),
            'fecha': col('dia y fecha de la call'),
            'fuente': col('fuente'),
            'contexto': col('situacion del lead\n(contexto', 'situacion del lead (contexto'),
            'presento': col('se presento?'),
            'califica': col('calificaba?'),
            'situacion': col('situacion del lead'),
            'pago_call': col('¿cuanto pago en llamada'),
            'sensacion': col('¿como te sentiste'),
            'manychat': col('url de manychat'),
            'fathom': col('link llamada de fathom'),
            'reagenda': col('fecha y hora de reagenda'),
            'pago_seg': col('¿cuanto pago en seguimiento', '¿cuanto pago antes'),
            'presento_re': col('se presento a la reagenda'),
            'closer': col('closer'),
        }
        # "Situación del Lead" aparece dos veces: la de contexto y la real
        if cols['contexto'] is not None and cols['situacion'] == cols['contexto']:
            posteriores = [i for h, i in indice.items()
                           if h.startswith('situacion del lead') and i != cols['contexto']]
            cols['situacion'] = posteriores[0] if posteriores else cols['situacion']

        for fila in ws.iter_rows(min_row=2, values_only=True):
            nombre = texto(fila[c_nombre]) if c_nombre < len(fila) else ''
            if not nombre:
                continue

            def v(clave):
                i = cols.get(clave)
                return fila[i] if i is not None and i < len(fila) else None

            situacion = texto(v('situacion'))
            pago_call = num(v('pago_call')) or 0
            pago_seg = num(v('pago_seg')) or 0
            estado = estado_de(situacion, texto(v('presento')))
            if pago_call or pago_seg:
                estado = 'Ganado'
            fecha = fecha_de_hoja(v('fecha'), mes)
            fuente = texto(v('fuente'))

            salida.append({
                'id': 'rec_lead_' + re.sub(r'[^a-z0-9]+', '', sinacentos(nombre))[:18],
                'nombre': nombre,
                'telefono': telefono(v('tel')),
                'email': '', 'instagram': '',
                'origen': fuente,
                'estado': estado,
                'setter': 'Facundo' if fuente.lower().startswith('setter') else '',
                'closer': texto(v('closer')) or 'Gabo',
                'fecha_contacto': fecha,
                'fecha_llamada': (fecha + 'T12:00') if fecha else '',
                'oferta': '', 'precio': None,
                'cash_collected': (pago_call + pago_seg) or None,
                'probabilidad': {'Ganado': 100, 'Perdido': 0, 'Esperando pago': 80,
                                 'Seguimiento': 50, 'No Show': 15}.get(estado, 40),
                'interes': None,
                'se_presento': 'Sí' if sinacentos(v('presento')) == 'si' else ('No' if v('presento') else ''),
                'calificaba': 'Sí' if sinacentos(v('califica')) == 'si' else ('No' if v('califica') else ''),
                'pago_llamada': pago_call or None,
                'pago_seguimiento': pago_seg or None,
                'reagenda': iso(v('reagenda')),
                'presento_reagenda': texto(v('presento_re')),
                'contexto': texto(v('contexto')),
                'sensacion': texto(v('sensacion')),
                'manychat': texto(v('manychat')),
                'grabacion': texto(v('fathom')),
                'proximo_paso': '', 'fecha_proximo_paso': '', 'objecion': '', 'notas': ''
            })


# ------------------------------------------------------------------ setter

def leer_setter(path, dias, config):
    wb = openpyxl.load_workbook(path, data_only=True)
    if 'Config' in wb.sheetnames:
        for fila in wb['Config'].iter_rows(values_only=True):
            etiqueta = sinacentos(fila[0] or '')
            if etiqueta == 'setter' and fila[1]:
                config['setter'] = texto(fila[1])
            if etiqueta.startswith('objetivo minimo') and fila[1] is not None:
                config['min'] = float(fila[1])
            if etiqueta.startswith('objetivo maximo') and fila[1] is not None:
                config['max'] = float(fila[1])

    for ws in wb.worksheets:
        partes = ws.title.split()
        if len(partes) != 2 or sinacentos(partes[0]) not in MESES:
            continue
        for fila in ws.iter_rows(min_row=4, values_only=True):
            fecha = iso(fila[0])
            if not fecha:
                continue
            valores = [num(fila[i]) for i in (2, 3, 4, 5, 6)]
            if not any(v for v in valores):
                continue
            dias.append({
                'id': 'rec_dia_' + fecha,
                'fecha': fecha,
                'setter': config.get('setter', 'Facundo'),
                'conversaciones': valores[0] or 0,
                'outbound': valores[1] or 0,
                'seguimientos': valores[2] or 0,
                'agendas': valores[3] or 0,
                'pendientes': valores[4],
                'notas': texto(fila[9]) if len(fila) > 9 else ''
            })


# ------------------------------------------------------------------ alumnos

def leer_alumnos(path, alumnos, pagos):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb['Alumnos']
    cabeceras = [texto(c.value) for c in ws[3]]
    indice = {sinacentos(h): i for i, h in enumerate(cabeceras) if h}

    def col(clave):
        for h, i in indice.items():
            if h.startswith(sinacentos(clave)):
                return i
        return None

    campos = {k: col(k) for k in [
        'nombre completo', 'email', 'telefono', 'programa', 'fecha de ingreso',
        'duracion', 'fecha de finalizacion', 'estado', 'modalidad de pago',
        'precio total', 'total pagado', 'saldo pendiente', 'proxima cuota',
        'comprobantes de pago', 'resultados', '¿caso de exito?']}

    for fila in ws.iter_rows(min_row=4, values_only=True):
        def v(clave):
            i = campos.get(clave)
            return fila[i] if i is not None and i < len(fila) else None

        nombre = texto(v('nombre completo'))
        if not nombre:
            continue
        programa = texto(v('programa'))
        if sinacentos(programa) == 'downsell':
            programa = 'Downsell'
        ingreso = iso(v('fecha de ingreso'))
        pagado = num(v('total pagado')) or 0
        precio = num(v('precio total')) or 0
        idr = 'rec_alu_' + re.sub(r'[^a-z0-9]+', '', sinacentos(nombre))[:18]

        alumnos.append({
            'id': idr, 'nombre': nombre,
            'email': texto(v('email')), 'telefono': telefono(v('telefono')),
            'programa': programa, 'estado': texto(v('estado')) or 'Activo',
            'fecha_ingreso': ingreso, 'duracion': num(v('duracion')) or 90,
            'fecha_fin': iso(v('fecha de finalizacion')), 'dias_restantes': None,
            'modalidad': texto(v('modalidad de pago')),
            'precio_total': precio, 'total_pagado': pagado, 'saldo': precio - pagado,
            'proxima_cuota': iso(v('proxima cuota')),
            'lead': '', 'comprobantes': texto(v('comprobantes de pago')),
            'resultados': texto(v('resultados')),
            'caso_exito': sinacentos(v('¿caso de exito?')) in ('si', 'sí', 'true', 'x')
        })

        if pagado:
            plan = sinacentos(v('modalidad de pago')).startswith('plan')
            pagos.append({
                'id': 'rec_pago_' + idr[8:],
                'concepto': ('Cuota' if plan else 'Pago total') + ' — ' + nombre,
                'alumno': idr, 'lead': '', 'fecha': ingreso, 'cuota': 1 if plan else None,
                'monto': pagado, 'tipo': 'Cuota' if plan else 'Pago total',
                'metodo': '', 'closer': '', 'factura': '',
                'notas': 'Migrado de Gestión de Alumnos.'
            })


# ------------------------------------------------------------------ salida

def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)

    leads, dias, alumnos, pagos = [], [], [], []
    config = {'setter': 'Facundo', 'min': 0.08, 'max': 0.12}

    for path in sys.argv[1:]:
        wb = openpyxl.load_workbook(path, data_only=True)
        hojas = wb.sheetnames
        if 'Alumnos' in hojas:
            leer_alumnos(path, alumnos, pagos)
        elif 'Config' in hojas:
            leer_setter(path, dias, config)
        else:
            leer_closer(path, leads, config)

    # Vincular alumnos con su lead por nombre
    por_nombre = {sinacentos(l['nombre']): l['id'] for l in leads}
    for a in alumnos:
        lid = por_nombre.get(sinacentos(a['nombre']))
        if lid:
            a['lead'] = lid
            for p in pagos:
                if p['alumno'] == a['id']:
                    p['lead'] = lid

    equipo = [
        ('Mariano', 'Dueño', 'mariano@alphaecommerce.com', 0, 0),
        ('Admin', 'Admin', '', 0, 0),
        (config['setter'], 'Setter', '', 0, 5),
        ('Gabo', 'Closer', '', 60000, 10),
    ]
    equipo = [{
        'id': 'rec_eq_' + sinacentos(n).replace(' ', ''), 'nombre': n, 'rol': r,
        'email': e, 'meta_cash': m, 'comision': c, 'activo': True,
        'clave': sha256(sinacentos(n).split()[0]), 'claveInicial': True
    } for n, r, e, m, c in equipo]

    # --- Tareas de arranque: salen de lo que los trackers dejan pendiente ---
    tareas = []

    def tarea(titulo, asignados, prioridad, area, detalle=''):
        tareas.append({
            'id': 'rec_tar_%02d' % (len(tareas) + 1), 'titulo': titulo,
            'asignados': asignados, 'estado': 'Pendiente', 'hecha': False,
            'prioridad': prioridad, 'area': area, 'vence': '', 'lead': '',
            'detalle': detalle, 'creada_por': 'Admin'
        })

    for a in sorted(alumnos, key=lambda x: -x['saldo']):
        if a['saldo'] > 0:
            tarea('Cobrar saldo de %s (%s USD)' % (a['nombre'], a['saldo']),
                  ['Admin'], 'Alta', 'Administración',
                  'Programa %s. Próxima cuota: %s' % (a['programa'], a['proxima_cuota'] or 'sin definir'))

    # Plata que figura cobrada en el CRM del closer pero no en Gestión de Alumnos
    pagado_por_lead = {}
    for p in pagos:
        if p['lead']:
            pagado_por_lead[p['lead']] = pagado_por_lead.get(p['lead'], 0) + p['monto']
    for l in leads:
        cobrado_crm = l.get('cash_collected') or 0
        if cobrado_crm and cobrado_crm > pagado_por_lead.get(l['id'], 0):
            tarea('Conciliar el pago de %s (%s USD)' % (l['nombre'], cobrado_crm),
                  ['Admin', 'Gabo'], 'Alta', 'Administración',
                  'Figura cobrado en el CRM del closer pero no aparece en Gestión de Alumnos.')

    faltan_datos = [a['nombre'] for a in alumnos if not a['email'] or not a['telefono']]
    if faltan_datos:
        tarea('Completar email y teléfono de los alumnos', ['Admin'], 'Media', 'Operaciones',
              'Faltan datos de: ' + ', '.join(faltan_datos))

    tarea('Cargar los links de los recursos (scripts, contratos, formación)',
          ['Mariano'], 'Media', 'Operaciones',
          'Están creados en la tabla Recursos, sólo falta pegar cada link.')

    recursos = [{
        'id': 'rec_res_%02d' % (i + 1), 'titulo': t, 'categoria': c, 'para': para,
        'url': '', 'descripcion': d, 'destacado': destacado
    } for i, (t, c, para, d, destacado) in enumerate([
        ('Script de setting (DM a llamada)', 'Scripts', 'Setter', 'Pegá acá el link al documento.', True),
        ('Script de cierre', 'Scripts', 'Closer', 'Pegá acá el link al documento.', True),
        ('Manejo de objeciones', 'Formación', 'Closer', 'Precio, tiempo, pareja, confianza.', True),
        ('Contrato de servicio', 'Contratos', 'Closer', 'Modelo para enviar al cerrar.', False),
        ('Carpeta de comprobantes de pago', 'Otros', 'Admin', 'Drive con los comprobantes.', False),
        ('Landing de Alpha Ecommerce', 'Otros', 'Todos', 'La página que ven los leads.', False),
    ])]
    recursos[5]['url'] = '../index.html'

    meta_cash = config.get('meta_cash') or 0
    hoy = datetime.date.today()
    metas = []
    for i in range(-1, 3):
        total = (hoy.year * 12 + hoy.month - 1) + i
        mes = datetime.date(total // 12, total % 12 + 1, 1)
        metas.append({
            'id': 'rec_meta_' + mes.strftime('%Y%m'), 'mes': mes.strftime('%Y-%m'),
            'meta_cash': meta_cash, 'meta_leads': None, 'meta_agendas': None,
            'meta_cierres': None,
            'nota': 'KPI mensual del equipo de closers' if meta_cash else ''
        })

    datos = {
        'equipo': equipo, 'leads': leads, 'alumnos': alumnos, 'pagos': pagos,
        'setter_dia': dias, 'actividades': [], 'contenido': [], 'tareas': tareas,
        'recursos': recursos, 'metas': metas
    }

    js = ('/* ===================================================================\n'
          '   Alpha CRM — datos migrados de los trackers de Excel de Alpha.\n'
          '   Generado por tools/importar-trackers.py — no editar a mano:\n'
          '   una vez cargado en la app, los cambios se hacen desde la app.\n'
          '   =================================================================== */\n'
          '(function (AE) {\n'
          "  'use strict';\n"
          '  AE.datosIniciales = ' + json.dumps(datos, ensure_ascii=False, indent=2) + ';\n'
          '  AE.datosIniciales.objetivos = ' +
          json.dumps({'tasaAgendaMin': config['min'], 'tasaAgendaMax': config['max']}) + ';\n'
          '})(window.AE);\n')

    with open('app/assets/js/datos-alpha.js', 'w', encoding='utf-8') as f:
        f.write(js)

    print('Escrito app/assets/js/datos-alpha.js')
    for k, v in datos.items():
        if v:
            print('  %-12s %d' % (k, len(v)))


if __name__ == '__main__':
    main()
